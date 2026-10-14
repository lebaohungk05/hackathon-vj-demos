"""Block engine for building styled Excel workbooks - the skill's own copy.

Drop this next to your content modules and import `build_workbook` from it. The only
third-party requirement is `openpyxl`; everything else is stdlib, so it runs in any repo
without carrying anything else along.

    from xlsx_style import build_workbook

    build_workbook([("Sheet 1", blocks)], Path("out.xlsx"),
                   image_dir=Path("screenshots"),
                   callout_labels={"ok": "CONCLUSION:", "warn": "WARNING:"})

`references/formats.md` lists the block types and the problems this engine already solves -
read that before deciding to write your own.

## Provenance

Copied from `scripts/docs/xlsx_style.py` in the fabric-model-analyzer repository, which is
where it is exercised daily against four workbooks in three languages. **That copy is the
original.** If you improve this one, port the change back; if you find the two have drifted,
trust the repository's.

The alternative was a skill that names a file in one particular repo and is therefore useless
in every other one. A copy that can drift is the lesser problem, and naming which side is
authoritative is what keeps it manageable.

## Style

Calibri body, Consolas for code, a light grey table header, thin `#D5DBE1` borders, no row
banding, and two soft callout tints. Column A is a narrow gutter; content lives in B..E.

Content is supplied as a list of blocks so one renderer serves every language edition of a
document - which is what makes a figure impossible to disagree between them.
"""
from __future__ import annotations

import unicodedata
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

COLUMN_WIDTHS = {"A": 3, "B": 40, "C": 46, "D": 34, "E": 40}
FIRST_COL = 2   # B
LAST_COL = 5    # E

TITLE_BLUE = "FF1B5E9C"
INK = "FF1F2933"
MUTED = "FF6B7280"
SECTION_BAND = "FFE8EEF4"
TH_BAND = "FFF2F4F7"
CODE_BAND = "FFF6F8FA"
OK_BAND = "FFEDF7ED"
WARN_BAND = "FFFFF8E1"
BORDER_GREY = "FFD5DBE1"

# Tab colours, and the matching band for a heading on the contents sheet. A workbook that
# holds two independent methods needs the tab strip to say which is which; a number alone
# does not, because 1.2 and 2.1 sit next to each other and read as one sequence.
TAB_SUMMARY = "FF9AA5B1"
TAB_METHOD1 = "FF4C9A4C"
TAB_METHOD2 = "FF1B5E9C"
BAND_METHOD1 = "FFE6F2E6"
BAND_METHOD2 = "FFE1ECF7"

_thin = Side(style="thin", color=BORDER_GREY)
CELL_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

FONT_TITLE = Font(name="Calibri", size=15, bold=True, color=TITLE_BLUE)
FONT_SUBTITLE = Font(name="Calibri", size=10, color=MUTED)
FONT_SECTION = Font(name="Calibri", size=12, bold=True, color=INK)
FONT_TH = Font(name="Calibri", size=10, bold=True, color=INK)
FONT_BODY = Font(name="Calibri", size=10, color=INK)
FONT_CODE = Font(name="Consolas", size=9, color=INK)
FONT_CAPTION = Font(name="Calibri", size=9, italic=True, color=MUTED)
FONT_LINK = Font(name="Calibri", size=10, color=TITLE_BLUE, underline="single")
# The "where was this taken" line above a screenshot. Small and bold rather than loud: it is
# orientation, not a heading.
FONT_WHERE = Font(name="Calibri", size=9, bold=True, color=MUTED)

# indent=1 keeps text off the cell border, which is what makes a bordered table
# readable. The template leaves it at 0 and reads cramped.
WRAP_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
WRAP_MID = Alignment(vertical="center", wrap_text=True, indent=1)
CODE_ALIGN = Alignment(horizontal="left", vertical="center", indent=1)
BULLET_ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)

CALLOUT_FILLS = {"ok": OK_BAND, "warn": WARN_BAND, "info": TH_BAND}

# Characters that fit per row when text is merged across N columns, from the
# B..E widths above. Used to estimate row height so wrapped text is not clipped.
_CHARS_PER_LINE = {1: 36, 2: 80, 3: 114, 4: 152}
# A little over the 12.75 pt a 10 pt line strictly needs, so wrapped text has some
# air between lines instead of sitting flush.
_LINE_HEIGHT = 14.4
_ROW_PADDING = 5.0

# Columns B..E come to 160 character units, roughly 1120 px, so keep images just
# inside that so none spills past column E.
MAX_IMAGE_WIDTH_PX = 1080
MAX_IMAGE_HEIGHT_PX = 560
# Pixels of image height covered by one reserved spacer row.
_IMAGE_ROW_PX = 20


def _display_width(text: str) -> int:
    """Width in character units, counting CJK as two.

    _CHARS_PER_LINE is calibrated on Latin text. Japanese and Chinese characters occupy two
    of those units, so counting them as one told the Japanese edition a line held twice the
    text it does, and every row that wrapped came out one line short and clipped.
    """
    return sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in text)


def _estimate_height(text: Any, span: int, min_height: float = 17.0) -> float:
    per_line = _CHARS_PER_LINE.get(min(max(span, 1), 4), 152)
    lines = 0
    for paragraph in str(text).split("\n"):
        lines += max(1, -(-_display_width(paragraph) // per_line))
    return max(min_height, lines * _LINE_HEIGHT + _ROW_PADDING)


def _link_of(value: Any) -> tuple[Any, str | None]:
    """Unpack a ``(text, url)`` cell into its parts.

    Table cells may carry a hyperlink by passing a two-item sequence whose second
    item looks like a URL; anything else is returned as a plain value.
    """
    if (isinstance(value, (tuple, list)) and len(value) == 2
            and isinstance(value[1], str)
            and value[1].startswith(("http://", "https://", "mailto:"))):
        return value[0], value[1]
    return value, None


class SheetWriter:
    """Appends styled blocks down a worksheet, tracking the current row."""

    def __init__(self, ws: Worksheet, image_dir: Path | None = None,
                 callout_labels: dict[str, str] | None = None) -> None:
        self.ws = ws
        self.row = 1
        self.image_dir = image_dir
        # The template relies on colour alone; a short word keeps the meaning
        # readable in print and for anyone who cannot distinguish the tints.
        self.callout_labels = callout_labels or {}
        for col, width in COLUMN_WIDTHS.items():
            ws.column_dimensions[col].width = width
        ws.sheet_view.showGridLines = False

    # -- primitives ------------------------------------------------------

    def _span(self, text: Any, first: int, last: int, font: Font,
              alignment: Alignment, fill: str | None = None,
              border: bool = False, height: float | None = None) -> None:
        ws = self.ws
        cell = ws.cell(row=self.row, column=first, value=text)
        cell.font = font
        cell.alignment = alignment
        if fill:
            patterned = PatternFill("solid", fgColor=fill)
            for col in range(first, last + 1):
                ws.cell(row=self.row, column=col).fill = patterned
        if border:
            for col in range(first, last + 1):
                ws.cell(row=self.row, column=col).border = CELL_BORDER
        if last > first:
            ws.merge_cells(start_row=self.row, start_column=first,
                           end_row=self.row, end_column=last)
        ws.row_dimensions[self.row].height = height or 15.0
        self.row += 1

    def blank(self, count: int = 1) -> None:
        self.row += count

    # -- blocks ----------------------------------------------------------

    def title(self, text: str) -> None:
        self._span(text, FIRST_COL, LAST_COL, FONT_TITLE, WRAP_MID, height=28)

    def subtitle(self, text: str) -> None:
        # No trailing blank: whatever follows is a section, which adds its own.
        self._span(text, FIRST_COL, LAST_COL, FONT_SUBTITLE, WRAP_TOP,
                   height=_estimate_height(text, 4, 28))

    def section(self, text: str, band: str | None = None) -> None:
        # Two rows of air above a section heading, one below, so the sheet reads as
        # blocks rather than one continuous wall.
        #
        # `band` overrides the default tint. The contents sheet uses it to give each of the
        # two methods the colour of its own tab group, so the reader can match a heading to
        # a tab without reading either.
        self.blank(2)
        self._span(text, FIRST_COL, LAST_COL, FONT_SECTION, WRAP_MID,
                   fill=band or SECTION_BAND, height=24)
        self.blank()

    def note(self, text: str) -> None:
        self._span(text, FIRST_COL, LAST_COL, FONT_BODY, WRAP_TOP,
                   height=_estimate_height(text, 4))
        self.blank()

    def bullet(self, text: str) -> None:
        # Bullets stay tight - a list should read as a list.
        body = f"-  {text}"
        self._span(body, FIRST_COL, LAST_COL, FONT_BODY, BULLET_ALIGN,
                   height=_estimate_height(body, 4))

    def callout(self, kind: str, text: str) -> None:
        label = self.callout_labels.get(kind)
        body = f"{label}  {text}" if label else str(text)
        self._span(body, FIRST_COL, LAST_COL, FONT_BODY, WRAP_TOP,
                   fill=CALLOUT_FILLS.get(kind, TH_BAND),
                   height=_estimate_height(body, 4, 24))
        self.blank()

    def code(self, lines: Iterable[str]) -> None:
        for line in lines:
            self._span(line if str(line).strip() else " ", FIRST_COL, LAST_COL,
                       FONT_CODE, CODE_ALIGN, fill=CODE_BAND, height=13.8)
        self.blank()

    def caption(self, text: str) -> None:
        self._span(text, FIRST_COL, LAST_COL, FONT_CAPTION, WRAP_TOP,
                   height=_estimate_height(text, 4, 16))

    def image(self, filename: str, caption: str, where: str | None = None) -> None:
        """Insert a screenshot, or a placeholder row when the file is absent.

        `where` names the product the capture was taken in, on a line above it. Screenshots are
        cropped to the part that matters, which routinely crops away the one thing that says
        which portal you are looking at - the Entra admin center and the Fabric admin portal
        are a pane of settings either way.

        The screenshots differ wildly in shape - wide portal captures around
        1830x820, a tall narrow folder tree at 386x787 - so the image is fitted
        inside a box rather than scaled to a fixed height. Scaling by height alone
        would shrink the tall one to a sliver and overflow the wide one past
        column E.
        """
        if where:
            self._span(where, FIRST_COL, LAST_COL, FONT_WHERE, WRAP_MID,
                       fill=TH_BAND, height=16)

        path = (self.image_dir / filename) if self.image_dir else None
        if not (path and path.exists()):
            self._span(" ", FIRST_COL, LAST_COL, FONT_CAPTION, WRAP_TOP,
                       fill=CODE_BAND, border=True, height=18)
            self.caption(caption)
            self.blank()
            return

        img = XLImage(str(path))
        if img.width and img.height:
            scale = min(MAX_IMAGE_WIDTH_PX / img.width,
                        MAX_IMAGE_HEIGHT_PX / img.height,
                        1.0)
            img.width = int(img.width * scale)
            img.height = int(img.height * scale)

        self.ws.add_image(img, f"{get_column_letter(FIRST_COL)}{self.row}")

        # Reserve real vertical space, otherwise the image floats over the caption
        # and whatever follows it.
        rows = max(1, -(-int(img.height or 0) // _IMAGE_ROW_PX))
        for offset in range(rows):
            self.ws.row_dimensions[self.row + offset].height = _IMAGE_ROW_PX * 0.75
        self.row += rows

        self.caption(caption)
        self.blank()

    def table(self, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
        """A bordered table starting at column B.

        The final column stretches to E when the table has fewer than four
        columns, matching the template's B/C/D:E shape. Rows are not banded -
        the template relies on the borders alone.
        """
        spans = self._column_spans(len(headers))
        self._row_of_cells(headers, spans, FONT_TH, WRAP_MID, fill=TH_BAND)
        for row in rows:
            self._row_of_cells(row, spans, FONT_BODY, WRAP_TOP)
        self.blank()

    def _column_spans(self, ncols: int) -> list[tuple[int, int]]:
        spans: list[tuple[int, int]] = []
        col = FIRST_COL
        for index in range(ncols):
            if index == ncols - 1:
                spans.append((col, LAST_COL))
            else:
                spans.append((col, col))
                col += 1
        return spans

    def _row_of_cells(self, values: Sequence[Any], spans: Sequence[tuple[int, int]],
                      font: Font, alignment: Alignment, fill: str | None = None) -> None:
        ws = self.ws
        patterned = PatternFill("solid", fgColor=fill) if fill else None
        tallest = 15.0

        for index, (first, last) in enumerate(spans):
            raw = values[index] if index < len(values) else None
            value, url = _link_of(raw)
            cell = ws.cell(row=self.row, column=first, value=value)
            cell.font = FONT_LINK if url else font
            cell.alignment = alignment
            if url:
                cell.hyperlink = url
            for col in range(first, last + 1):
                target = ws.cell(row=self.row, column=col)
                target.border = CELL_BORDER
                if patterned:
                    target.fill = patterned
            if last > first:
                ws.merge_cells(start_row=self.row, start_column=first,
                               end_row=self.row, end_column=last)
            if value:
                tallest = max(tallest, _estimate_height(value, last - first + 1))

        ws.row_dimensions[self.row].height = tallest
        self.row += 1


BLOCK_DISPATCH = {
    "title": lambda w, b: w.title(b[1]),
    "subtitle": lambda w, b: w.subtitle(b[1]),
    "section": lambda w, b: w.section(b[1], b[2] if len(b) > 2 else None),
    "note": lambda w, b: w.note(b[1]),
    "bullet": lambda w, b: w.bullet(b[1]),
    "ok": lambda w, b: w.callout("ok", b[1]),
    "warn": lambda w, b: w.callout("warn", b[1]),
    "info": lambda w, b: w.callout("info", b[1]),
    "code": lambda w, b: w.code(b[1]),
    "caption": lambda w, b: w.caption(b[1]),
    "image": lambda w, b: w.image(b[1], b[2], b[3] if len(b) > 3 else None),
    "table": lambda w, b: w.table(b[1], b[2]),
    "blank": lambda w, b: w.blank(b[1] if len(b) > 1 else 1),
}


def build_workbook(sheets: Sequence[tuple[str, Sequence[tuple]]],
                   output: Path,
                   image_dir: Path | None = None,
                   callout_labels: dict[str, str] | None = None) -> Path:
    """Render ``sheets`` to ``output``.

    Each entry is ``(sheet_name, blocks)``, growing to
    ``(sheet_name, blocks, tab_colour, hidden)`` as needed - a tinted tab, and a sheet that is
    built but not shown. The two-item form is what the older workbooks pass, and it stays
    valid, so a document only pays for the parts it uses.
    """
    wb = Workbook()
    wb.remove(wb.active)

    for entry in sheets:
        name, blocks = entry[0], entry[1]
        tab_colour = entry[2] if len(entry) > 2 else None
        hidden = entry[3] if len(entry) > 3 else False
        ws = wb.create_sheet(title=name[:31])
        if tab_colour:
            ws.sheet_properties.tabColor = tab_colour
        # Hidden rather than deleted: a sheet that is out of scope this week is usually back in
        # scope later, and rebuilding it from a commit is slower than unhiding it.
        if hidden:
            ws.sheet_state = "hidden"
        writer = SheetWriter(ws, image_dir=image_dir, callout_labels=callout_labels)
        for block in blocks:
            handler = BLOCK_DISPATCH.get(block[0])
            if handler is None:
                raise ValueError(f"Unknown block type {block[0]!r} on sheet {name!r}")
            handler(writer, block)
        # Keep the title and subtitle in view while scrolling a long sheet.
        ws.freeze_panes = "A4"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = False

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
